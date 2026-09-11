from pptx import Presentation

def generate_slides():
    prs = Presentation(r'E:\CAP\Documentation\G1-PPT-Review-I.pptx')

    # Delete existing slides to leave a clean deck with just the template
    xml_slides = prs.slides._sldIdLst  
    slides = list(xml_slides)
    for s in slides:
        xml_slides.remove(s)

    layout = prs.slide_layouts[1] # Title and Content Layout

    # --- SLIDE 1: Implementation Progress ---
    slide1 = prs.slides.add_slide(layout)
    slide1.shapes.title.text = "Implementation Progress (>50% Completed)"
    tf1 = slide1.shapes.placeholders[1].text_frame
    tf1.text = "Machine Learning Grid Search is fully built and running (bakery_model.py)."
    p = tf1.add_paragraph()
    p.text = "Cross-Validation pipeline is active and calculating baseline RMSE metrics."
    p = tf1.add_paragraph()
    p.text = "Flutter 'Warm Tech' UI is successfully scaffolded via Stitch MCP."
    p = tf1.add_paragraph()
    p.text = "FastAPI bridge architecture is finalized and ready for ML integration."

    # --- SLIDE 2: Working Demonstration ---
    slide2 = prs.slides.add_slide(layout)
    slide2.shapes.title.text = "Working Demonstration & Initial Results"
    tf2 = slide2.shapes.placeholders[1].text_frame
    tf2.text = "[ Paste screenshot of Python terminal showing Prophet Grid Search / RMSE here ]"
    p = tf2.add_paragraph()
    p.text = "[ Paste screenshot of Flutter mobile app UI (Dashboard) here ]"

    # --- SLIDE 3: Remaining Work, Future Extensions & Status ---
    # Combining Points 4 and 5 into one slide as requested
    slide3 = prs.slides.add_slide(layout)
    slide3.shapes.title.text = "Project Status, Remaining Work & Future Extensions"
    tf3 = slide3.shapes.placeholders[1].text_frame
    
    tf3.text = "Gantt Chart Progress: Project is strictly on schedule per the Review 1 timeline."
    
    p = tf3.add_paragraph()
    p.text = "Key Challenges Overcome:"
    p.level = 0
    p2 = tf3.add_paragraph()
    p2.text = "Addressed SME data sparsity via XGBoost hybrid and bypassed weather API rate limits using caching."
    p2.level = 1
    
    p3 = tf3.add_paragraph()
    p3.text = "Remaining Work (Review 3):"
    p3.level = 0
    p4 = tf3.add_paragraph()
    p4.text = "FastAPI integration, Waste Action Engine intraday logic, and Indian holdout dataset testing."
    p4.level = 1
    
    p5 = tf3.add_paragraph()
    p5.text = "Future Extensions (Beyond Scope):"
    p5.level = 0
    p6 = tf3.add_paragraph()
    p6.text = "Multi-location chain forecasting, live POS hardware integration, and interactive LLM chatbot."
    p6.level = 1

    out_path = r'E:\CAP\Documentation\G2-PPT-Review-II-Targeted.pptx'
    prs.save(out_path)
    print(f"Saved successfully to {out_path}")

if __name__ == "__main__":
    generate_slides()
