from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def design_professional_slides():
    prs = Presentation(r'E:\CAP\Documentation\G1-PPT-Review-I.pptx')

    # Keep the template, delete all existing slides
    xml_slides = prs.slides._sldIdLst  
    slides = list(xml_slides)
    for s in slides:
        xml_slides.remove(s)

    title_only_layout = prs.slide_layouts[5] # Layout with just the Title

    # ---------------------------------------------------------
    # SLIDE 1: Implementation Progress
    # ---------------------------------------------------------
    s1 = prs.slides.add_slide(title_only_layout)
    s1.shapes.title.text = "Implementation Progress (>50% Completed)"

    # Add 3 professional boxes for ML, Frontend, Backend
    colors = [RGBColor(27, 67, 50), RGBColor(255, 183, 2), RGBColor(33, 150, 243)] # Deep Green, Amber, Blue
    titles = ["Machine Learning Core", "Frontend Mobile UI", "Backend API Bridge"]
    texts = [
        "• Grid Search fully built and optimized (bakery_model.py).\n• Cross-Validation pipeline active and calculating baseline RMSE metrics.",
        "• Flutter 'Warm Tech' UI successfully scaffolded via Stitch.\n• Dashboard and input screens designed for non-technical users.",
        "• FastAPI bridge architecture finalized.\n• Ready to connect ML Simulator to Flutter frontend."
    ]

    top = Inches(2.0)
    width = Inches(8.5)
    height = Inches(1.2)
    left = Inches(0.75)

    for i in range(3):
        # Create rounded rectangle
        shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top + Inches(i * 1.5), width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = colors[i]
        shape.line.color.rgb = RGBColor(255, 255, 255)
        shape.line.width = Pt(1.5)
        
        tf = shape.text_frame
        tf.word_wrap = True
        
        # Add Title Paragraph
        p1 = tf.paragraphs[0]
        p1.text = titles[i]
        p1.font.bold = True
        p1.font.size = Pt(18)
        p1.font.color.rgb = RGBColor(255, 255, 255)
        
        # Add Body Paragraph
        p2 = tf.add_paragraph()
        p2.text = texts[i]
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(255, 255, 255)

    # ---------------------------------------------------------
    # SLIDE 2: Working Demonstration & Initial Results
    # ---------------------------------------------------------
    s2 = prs.slides.add_slide(title_only_layout)
    s2.shapes.title.text = "Working Demonstration & Initial Results"

    # Add two dashed placeholder boxes for screenshots
    left1 = Inches(0.5)
    top1 = Inches(2.0)
    width1 = Inches(4.25)
    height1 = Inches(4.5)

    left2 = Inches(5.25)

    # Box 1 (Left)
    box1 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left1, top1, width1, height1)
    box1.fill.solid()
    box1.fill.fore_color.rgb = RGBColor(240, 240, 240) # Light grey
    box1.line.color.rgb = RGBColor(100, 100, 100)
    box1.line.dash_style = 7 # Dashed
    box1.line.width = Pt(2)
    tf1 = box1.text_frame
    tf1.text = "PASTE SCREENSHOT HERE\n\nPython Terminal Output\n(Prophet Grid Search & RMSE)"
    tf1.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf1.paragraphs[0].font.color.rgb = RGBColor(100, 100, 100)
    tf1.paragraphs[0].font.bold = True

    # Box 2 (Right)
    box2 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left2, top1, width1, height1)
    box2.fill.solid()
    box2.fill.fore_color.rgb = RGBColor(240, 240, 240)
    box2.line.color.rgb = RGBColor(100, 100, 100)
    box2.line.dash_style = 7
    box2.line.width = Pt(2)
    tf2 = box2.text_frame
    tf2.text = "PASTE SCREENSHOT HERE\n\nFlutter Mobile App UI\n(Dashboard Scaffold)"
    tf2.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf2.paragraphs[0].font.color.rgb = RGBColor(100, 100, 100)
    tf2.paragraphs[0].font.bold = True

    # ---------------------------------------------------------
    # SLIDE 3: Project Status & Future Extensions
    # ---------------------------------------------------------
    s3 = prs.slides.add_slide(title_only_layout)
    s3.shapes.title.text = "Project Status, Remaining Work & Future Extensions"

    # Status Banner (Top)
    banner = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.8), Inches(9.0), Inches(0.8))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(27, 67, 50)
    banner.line.color.rgb = RGBColor(255, 255, 255)
    tf_banner = banner.text_frame
    tf_banner.text = "GANTT CHART STATUS: strictly on schedule per Review 1 timeline."
    tf_banner.paragraphs[0].font.bold = True
    tf_banner.paragraphs[0].font.size = Pt(16)
    tf_banner.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    tf_banner.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Left Column: Remaining Work
    left_col = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.9), Inches(4.3), Inches(4.0))
    left_col.fill.solid()
    left_col.fill.fore_color.rgb = RGBColor(245, 245, 245)
    left_col.line.color.rgb = RGBColor(27, 67, 50)
    left_col.line.width = Pt(2)
    tf_l = left_col.text_frame
    tf_l.word_wrap = True
    p_l1 = tf_l.paragraphs[0]
    p_l1.text = "Remaining Work (Review 3)"
    p_l1.font.bold = True
    p_l1.font.size = Pt(18)
    p_l1.font.color.rgb = RGBColor(27, 67, 50)
    
    p_l2 = tf_l.add_paragraph()
    p_l2.text = "• API Integration: Connect FastAPI to Flutter frontend."
    p_l2.font.size = Pt(14)
    p_l2.font.color.rgb = RGBColor(0, 0, 0)
    
    p_l3 = tf_l.add_paragraph()
    p_l3.text = "• Waste Action Engine: Code intraday logic for markdowns."
    p_l3.font.size = Pt(14)
    p_l3.font.color.rgb = RGBColor(0, 0, 0)
    
    p_l4 = tf_l.add_paragraph()
    p_l4.text = "• Generalization Test: Evaluate on Indian holdout dataset."
    p_l4.font.size = Pt(14)
    p_l4.font.color.rgb = RGBColor(0, 0, 0)

    # Right Column: Future Extensions
    right_col = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.2), Inches(2.9), Inches(4.3), Inches(4.0))
    right_col.fill.solid()
    right_col.fill.fore_color.rgb = RGBColor(245, 245, 245)
    right_col.line.color.rgb = RGBColor(255, 183, 2)
    right_col.line.width = Pt(2)
    tf_r = right_col.text_frame
    tf_r.word_wrap = True
    p_r1 = tf_r.paragraphs[0]
    p_r1.text = "Future Extensions"
    p_r1.font.bold = True
    p_r1.font.size = Pt(18)
    p_r1.font.color.rgb = RGBColor(255, 183, 2)
    
    p_r2 = tf_r.add_paragraph()
    p_r2.text = "• Multi-location chain forecasting."
    p_r2.font.size = Pt(14)
    p_r2.font.color.rgb = RGBColor(0, 0, 0)
    
    p_r3 = tf_r.add_paragraph()
    p_r3.text = "• Live POS (Point of Sale) hardware integration."
    p_r3.font.size = Pt(14)
    p_r3.font.color.rgb = RGBColor(0, 0, 0)
    
    p_r4 = tf_r.add_paragraph()
    p_r4.text = "• Interactive LLM chatbot for dynamic questioning."
    p_r4.font.size = Pt(14)
    p_r4.font.color.rgb = RGBColor(0, 0, 0)

    out_path = r'E:\CAP\Documentation\G2-PPT-Review-II-Designed.pptx'
    prs.save(out_path)
    print(f"Professional presentation saved to {out_path}")

if __name__ == "__main__":
    design_professional_slides()
