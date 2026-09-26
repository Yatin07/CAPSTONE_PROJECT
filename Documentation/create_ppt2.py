from pptx import Presentation
from pptx.util import Pt
import os

# Load the base presentation to copy its layout and style
base_path = r'E:\CAP\Documentation\G1-PPT-Review-I.pptx'
prs = Presentation(base_path)

# Delete all slides except the first one (Title slide) to keep the template fresh
# Actually, let's just delete all slides to have a clean slate of new slides.
xml_slides = prs.slides._sldIdLst  
slides = list(xml_slides)
for s in slides:
    xml_slides.remove(s)

title_content_layout = prs.slide_layouts[1]

# Slide 1: Literature Review & Article Status
slide = prs.slides.add_slide(title_content_layout)
slide.shapes.title.text = "Literature Review & Article Status"
tf = slide.shapes.placeholders[1].text_frame
tf.text = "Comprehensive Review: Analyzed 25+ recent papers (2024–2026) across forecasting methods, SME data challenges, and AI technology adoption."
p = tf.add_paragraph()
p.text = "Key Gap Identified: Existing models stop at calculating error metrics (RMSE/MAPE) and fail to close the 'Forecast-to-Action' gap for non-technical users."
p = tf.add_paragraph()
p.text = "Review Article Status:"
p = tf.add_paragraph()
p.text = "Title: Demand Forecasting and Inventory Recommendation in Small Retail: A Review"
p.level = 1
p = tf.add_paragraph()
p.text = "Format: Drafted and formatted to IEEE conference standards (LocalDemand_LitReview_IEEE.tex)."
p.level = 1
p = tf.add_paragraph()
p.text = "Current Status: Manuscript drafted; currently undergoing final internal proofreading before submission."
p.level = 1

# Slide 2: Tools and Technologies Used
slide = prs.slides.add_slide(title_content_layout)
slide.shapes.title.text = "Development & Progress: Tools and Technologies"
tf = slide.shapes.placeholders[1].text_frame
tf.text = "Machine Learning Engine: Python (Prophet for trend/seasonality, XGBoost for residual correction, Optuna for hyperparameter tuning)."
p = tf.add_paragraph()
p.text = "Backend API: FastAPI (Python) – serves as the lightweight bridge between the heavy ML models and the mobile app."
p = tf.add_paragraph()
p.text = "Frontend Mobile App: Flutter (Dart) – chosen for cross-platform delivery with a 'Warm Tech' UI design system to reduce math anxiety for bakery owners."
p = tf.add_paragraph()
p.text = "Database & Authentication: Firebase / Supabase – for user profiles, shop roles, and historical sales logs."
p = tf.add_paragraph()
p.text = "Explainability Layer: LLM (OpenAI/Gemini API) – acts as a plain-language narrator for the forecasts."

# Slide 3: Implementation Progress
slide = prs.slides.add_slide(title_content_layout)
slide.shapes.title.text = "Development & Progress: Implementation Status"
tf = slide.shapes.placeholders[1].text_frame
tf.text = "Data Engineering: Cleaned and pre-processed historical daily sales data for proxy datasets (processed_bakery.csv)."
p = tf.add_paragraph()
p.text = "Machine Learning Core:"
p = tf.add_paragraph()
p.text = "Successfully implemented Grid Search optimization on the Prophet baseline (bakery_model.py)."
p.level = 1
p = tf.add_paragraph()
p.text = "Built the time-series cross-validation pipeline (initial 180 days, 30-day horizon) to establish benchmark RMSE."
p.level = 1
p = tf.add_paragraph()
p.text = "Frontend Design Framework: Finalized the 'Warm Tech' design system via Stitch MCP (Dashboard, Restock Input, and Waste Action Alert components mapped)."
p = tf.add_paragraph()
p.text = "Architecture Flow: Locked the end-to-end data pipeline: Flutter App -> FastAPI -> ML Simulator -> LLM Narrator -> Flutter."

# Slide 4: Working Demonstration
slide = prs.slides.add_slide(title_content_layout)
slide.shapes.title.text = "Development & Progress: Working Demonstration"
tf = slide.shapes.placeholders[1].text_frame
tf.text = "[Please insert your terminal screenshot of bakery_model.py output here]"
p = tf.add_paragraph()
p.text = "Caption: Prophet Grid Search executing on Bakery dataset, computing optimal changepoint priors and minimizing Cross-Validated RMSE."
p = tf.add_paragraph()
p.text = "[Please insert your Flutter Dashboard Stitch design mockup here]"
p = tf.add_paragraph()
p.text = "Caption: Main Dashboard concept showing plain-language LLM recommendations and item-level restock numbers."

# Slide 5: Remaining Work
slide = prs.slides.add_slide(title_content_layout)
slide.shapes.title.text = "Development & Progress: Remaining Work"
tf = slide.shapes.placeholders[1].text_frame
tf.text = "ML Engine Finalization: Implement the XGBoost residual correction layer over the Prophet baseline and code the 3-phase cold-start logic."
p = tf.add_paragraph()
p.text = "Waste Action Engine: Build the intraday logic to flag slow-moving items and recommend markdown percentages using price-elasticity data."
p = tf.add_paragraph()
p.text = "API & Integration: Wrap the Python ML pipeline into FastAPI endpoints (/forecast) and connect it live to the Flutter frontend."
p = tf.add_paragraph()
p.text = "Generalization Test: Evaluate the completed, tuned system on an isolated, untouched Indian food-retail dataset to test real-world transferability."

# Slide 6: Gantt Chart Progress
slide = prs.slides.add_slide(title_content_layout)
slide.shapes.title.text = "Project Management: Progress Compared with Review 1 Plan"
tf = slide.shapes.placeholders[1].text_frame
tf.text = "Month 1 (Weeks 1-5): Literature Review & Dataset Verification"
p = tf.add_paragraph()
p.text = "Status: 100% Completed. On Track."
p.level = 1
p = tf.add_paragraph()
p.text = "Month 2 (Weeks 6-10): Model Training & App Setup"
p = tf.add_paragraph()
p.text = "Status: In Progress. Currently optimizing Prophet baseline and transitioning to Flutter UI scaffolding."
p.level = 1
p = tf.add_paragraph()
p.text = "Month 3 (Weeks 11-15): API Integration, Testing & Documentation"
p = tf.add_paragraph()
p.text = "Status: Upcoming."
p.level = 1
p = tf.add_paragraph()
p.text = "Overall Health: The project is currently on schedule according to the Review 1 Gantt Chart, with no major delays in the critical path."

# Slide 7: Challenges Encountered
slide = prs.slides.add_slide(title_content_layout)
slide.shapes.title.text = "Project Management: Challenges Encountered"
tf = slide.shapes.placeholders[1].text_frame
tf.text = "SME Data Sparsity: Standard models struggle with zero-inflated, sparse data common in small cafes."
p = tf.add_paragraph()
p.text = "Solution: Architecting a hybrid approach where XGBoost specifically corrects Prophet's non-linear errors."
p.level = 1
p = tf.add_paragraph()
p.text = "LLM Hallucination Risk: Using LLMs to explain forecasts risked the AI inventing restock numbers."
p = tf.add_paragraph()
p.text = "Solution: Strictly decoupled the math from the text; the deterministic Pandas engine does the math, the LLM acts purely as a narrator."
p.level = 1
p = tf.add_paragraph()
p.text = "Weather API Limits: Encountered throttling on free-tier weather data during ML training."
p = tf.add_paragraph()
p.text = "Solution: Implemented caching and batched requests for historical regressor data to prevent timeouts."
p.level = 1

out_path = r'E:\CAP\Documentation\G2-PPT-Review-II-New-Slides.pptx'
prs.save(out_path)
print(f"Created {out_path} successfully!")
