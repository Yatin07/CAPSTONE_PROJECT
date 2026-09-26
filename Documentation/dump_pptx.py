import pptx
import sys

def dump_pptx(filepath, out_file):
    prs = pptx.Presentation(filepath)
    out_file.write(f"--- DUMPING {filepath} ---\n")
    for i, slide in enumerate(prs.slides):
        out_file.write(f"SLIDE {i+1}:\n")
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text:
                out_file.write(f"  TEXT: {repr(shape.text)}\n")
            if shape.has_table:
                out_file.write("  TABLE:\n")
                for row in shape.table.rows:
                    out_file.write("    " + " | ".join([cell.text_frame.text.replace("\n", " ") for cell in row.cells]) + "\n")

with open('dump_utf8.txt', 'w', encoding='utf-8') as f:
    dump_pptx("G1-PPT-Review-I.pptx", f)
    dump_pptx("G1-PPT-Review-I 2 1.pptx", f)
